from app.core.jd_parser import parse_jd
from app.core.resume_parser import extract_text, parse_resume


def test_parse_sample_resume():
    text = (__import__("pathlib").Path(__file__).parents[1] / "sample_data/resumes/asha_rao.txt").read_text()
    r = parse_resume(text)
    assert r["name"] == "Asha Rao"
    assert r["email"] == "asha.rao.sample@example.com"
    assert r["cgpa"] == 8.6
    assert len(r["projects"]) == 2 and len(r["experience"]) == 1
    assert "XGBoost" in r["skills"] and "FastAPI" in r["skills"]
    assert r["projects"][0]["tech"][:2] == ["Python", "XGBoost"]


def test_wrapped_bullets_and_date_lines_stay_inside_project():
    text = """Test User
test@example.com
Projects
Demo App | Python, Flask
2025
• Built a service that does something quite useful and it keeps
going onto the next line without a full stop
• Second bullet here.
Other App | Java
• One bullet only.
"""
    r = parse_resume(text)
    assert [p["title"] for p in r["projects"]] == ["Demo App", "Other App"]
    assert len(r["projects"][0]["bullets"]) == 2
    assert r["projects"][0]["dates"] == "2025"
    assert "onto the next line" in r["projects"][0]["bullets"][0]


def test_docx_and_txt_extraction(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Jane Roe")
    d.add_paragraph("Skills")
    d.add_paragraph("Python, SQL")
    p = tmp_path / "r.docx"
    d.save(p)
    assert "Python, SQL" in extract_text(p.read_bytes(), "r.docx")
    assert extract_text(b"hello world", "x.txt") == "hello world"


def test_jd_parser_splits_required_and_preferred(jd_text):
    jd = parse_jd(jd_text)
    assert "Python" in jd["required_skills"] and "SQL" in jd["required_skills"]
    assert "AWS" in jd["preferred_skills"] and "RAG" in jd["preferred_skills"]
    assert jd["level"] == "entry"
    assert not set(jd["required_skills"]) & set(jd["preferred_skills"])
