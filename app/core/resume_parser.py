"""Resume parser: PDF / DOCX / TXT  ->  structured dict.

The parser is rule-based so it works offline. When an LLM is configured the
caller may optionally refine the result (see `question_gen`/`llm`).
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Union

from .skills import extract_skills

SECTION_ALIASES = {
    "education": ["education", "academic background", "academics", "qualifications"],
    "skills": ["skills", "technical skills", "core skills", "skills & tools", "tools & technologies"],
    "projects": ["projects", "academic projects", "personal projects", "project experience"],
    "experience": ["experience", "work experience", "internships", "internship", "professional experience", "employment"],
    "publications": ["publications", "research", "papers"],
    "achievements": ["achievements & activities", "achievements", "activities", "achievements and activities",
                     "hackathons & competitions", "extracurricular", "awards", "certifications", "certificates"],
    "summary": ["summary", "profile", "objective", "about me", "career objective"],
}

_ALIAS_TO_SECTION = {a: s for s, al in SECTION_ALIASES.items() for a in al}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{8,}\d)")
URL_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:github\.com|linkedin\.com)/[\w./-]+", re.IGNORECASE)
BULLET_RE = re.compile(r"^\s*[•\-•●▪*▪]\s*")
DATE_RE = re.compile(
    r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s*\d{4}|\b(?:19|20)\d{2}\b)\s*[–\-—to]+\s*"
    r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s*\d{4}|\b(?:19|20)\d{2}\b|Present|Current)",
    re.IGNORECASE,
)


# --------------------------------------------------------------------------- text extraction
def extract_text(source: Union[str, Path, bytes, io.BytesIO], filename: str | None = None) -> str:
    """Extract plain text from a PDF, DOCX or TXT file (path, bytes or file-like)."""
    if isinstance(source, (str, Path)) and Path(str(source)).exists():
        path = Path(str(source))
        filename = filename or path.name
        data = path.read_bytes()
    elif isinstance(source, bytes):
        data = source
    elif hasattr(source, "read"):
        data = source.read()
    else:
        raise ValueError("Unsupported source for extract_text")

    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        return _pdf_text(data)
    if name.endswith(".docx") or data[:2] == b"PK":
        return _docx_text(data)
    return data.decode("utf-8", errors="ignore")


def _pdf_text(data: bytes) -> str:
    try:
        import pymupdf as fitz  # PyMuPDF (new import name)
    except ImportError:
        try:
            import fitz  # older PyMuPDF
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("PyMuPDF is required to read PDFs: pip install pymupdf") from exc
    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text("text") for page in doc)


def _docx_text(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


# --------------------------------------------------------------------------- parsing
def _normalise(text: str) -> str:
    text = text.replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _heading_key(line: str) -> str | None:
    clean = re.sub(r"[^a-z& ]", "", line.lower()).strip()
    if not clean or len(clean) > 40:
        return None
    # headings are short; allow trailing date text removed
    return _ALIAS_TO_SECTION.get(clean)


def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        key = _heading_key(line)
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _parse_name(header_lines: list[str]) -> str:
    for line in header_lines[:4]:
        if EMAIL_RE.search(line) or PHONE_RE.search(line) or URL_RE.search(line):
            continue
        words = line.split()
        if 1 < len(words) <= 5 and all(re.match(r"^[A-Za-z.'-]+$", w) for w in words):
            return line.strip()
    return header_lines[0].strip() if header_lines else ""


_DATE_LINE_RE = re.compile(
    r"^\W*(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s*)?(?:19|20)\d{2}"
    r"(?:\s*[–\-—to]+\s*(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s*)?(?:(?:19|20)\d{2}|Present|Current))?\W*$",
    re.IGNORECASE,
)


def _parse_projects(lines: list[str]) -> list[dict]:
    """A project = a title line followed by bullet lines.

    PDFs wrap long bullets onto the next line and put dates on their own line,
    so: date-only lines are attached to the current entry, and a non-bullet line
    that follows an unfinished bullet is treated as that bullet's continuation.
    """
    projects: list[dict] = []
    current: dict | None = None
    for line in lines:
        if _DATE_LINE_RE.match(line):
            if current is not None:
                current["dates"] = line.strip()
            continue
        if BULLET_RE.match(line):
            bullet = BULLET_RE.sub("", line).strip()
            if current is None:
                current = {"title": "Untitled project", "tech": [], "dates": "", "bullets": []}
                projects.append(current)
            current["bullets"].append(bullet)
            continue
        if current and current["bullets"] and "|" not in line:
            last = current["bullets"][-1]
            unfinished = not re.search(r"[.!?]$", last)
            if unfinished or line[:1].islower():
                current["bullets"][-1] = last + " " + line
                continue
        title = DATE_RE.sub("", line).strip(" |–-—")
        tech = ""
        if "|" in title:
            title, tech = [p.strip() for p in title.split("|", 1)]
        current = {"title": title, "tech": [t.strip() for t in re.split(r"[,/]", tech) if t.strip()],
                   "dates": "", "bullets": []}
        projects.append(current)
    for p in projects:
        blob = " ".join([p["title"], " ".join(p["tech"]), " ".join(p["bullets"])])
        p["skills"] = extract_skills(blob)
    return [p for p in projects if p["title"] or p["bullets"]]


def _parse_education(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    for line in lines:
        if BULLET_RE.match(line):
            continue
        cgpa = re.search(r"(?:CGPA|GPA)[:\s]*([0-9]+\.?[0-9]*)", line, re.IGNORECASE)
        pct = re.search(r"(\d{2}\.?\d*)\s*%", line)
        entries.append({
            "text": line,
            "cgpa": float(cgpa.group(1)) if cgpa else None,
            "percentage": float(pct.group(1)) if pct else None,
        })
    return entries


def parse_resume(text: str) -> dict:
    """Parse resume text into structured data."""
    text = _normalise(text)
    sections = split_sections(text)
    header = sections.get("header", [])
    header_blob = "\n".join(header)

    email = EMAIL_RE.search(text)
    phone = PHONE_RE.search(header_blob) or PHONE_RE.search(text)
    links = sorted({m.group(0) for m in URL_RE.finditer(text)})

    skills_section = "\n".join(sections.get("skills", []))
    projects = _parse_projects(sections.get("projects", []))
    experience = _parse_projects(sections.get("experience", []))
    education = _parse_education(sections.get("education", []))

    skills_from_section = extract_skills(skills_section)
    all_skills = extract_skills(text)

    bullets = [b for p in projects + experience for b in p["bullets"]]
    cgpas = [e["cgpa"] for e in education if e["cgpa"]]

    return {
        "name": _parse_name(header),
        "email": email.group(0) if email else "",
        "phone": phone.group(1).strip() if phone else "",
        "links": links,
        "summary": " ".join(sections.get("summary", [])),
        "education": education,
        "cgpa": cgpas[0] if cgpas else None,
        "skills": all_skills,
        "skills_listed": skills_from_section,
        "projects": projects,
        "experience": experience,
        "publications": sections.get("publications", []),
        "achievements": sections.get("achievements", []),
        "bullets": bullets,
        "sections_found": [k for k in sections if k != "header" and sections[k]],
        "raw_text": text,
        "word_count": len(text.split()),
    }
